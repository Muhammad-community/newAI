from flask import Flask, render_template, request, jsonify, session, redirect, url_for
import os, urllib.parse, requests as req, base64, io, json, sqlite3, hashlib, secrets
from functools import wraps
from datetime import datetime, timedelta

try:
    from dotenv import load_dotenv
    load_dotenv()
except: pass

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "novamind-stable-key-2025-xK9mP2qR")
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=30)

GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "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")
GROQ_URL     = "https://api.groq.com/openai/v1/chat/completions"

MODELS = {
    "llama-3.3-70b-versatile": {"name": "LLaMA 3.3 70B",  "desc": "Eng kuchli, universal",     "icon": "🦙"},
    "llama-3.1-8b-instant":    {"name": "LLaMA 3.1 8B",   "desc": "Eng tez model",             "icon": "⚡"},
    "llama3-70b-8192":         {"name": "LLaMA 3 70B",    "desc": "Kuchli va barqaror",        "icon": "🔥"},
    "gemma2-9b-it":            {"name": "Gemma 2 9B",     "desc": "Google'dan, aniq javoblar", "icon": "💎"},
}
FREE_MODEL = "llama-3.1-8b-instant"

PLANS = {
    "free":    {"name":"Free",    "price":0,     "icon":"🆓","models":False,"imagine":False,"slides":False},
    "pro":     {"name":"Pro",     "price":9.99,  "icon":"⭐","models":True, "imagine":False,"slides":False},
    "premium": {"name":"Premium", "price":19.99, "icon":"💎","models":True, "imagine":True, "slides":False},
    "ultra":   {"name":"Ultra",   "price":39.99, "icon":"🚀","models":True, "imagine":True, "slides":True},
}

# PAYME / CLICK ma'lumotlari — .env da sozlang
PAYME_ID      = os.environ.get("PAYME_MERCHANT_ID", "")
CLICK_ID      = os.environ.get("CLICK_MERCHANT_ID", "")
ADMIN_SECRET  = os.environ.get("ADMIN_SECRET", "novamind-admin-2025")

# ── DATABASE ──────────────────────────────────
DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "novamind.db")

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    return conn

def init_db():
    with get_db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS users (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            username   TEXT UNIQUE NOT NULL COLLATE NOCASE,
            email      TEXT UNIQUE NOT NULL COLLATE NOCASE,
            password   TEXT NOT NULL,
            plan       TEXT NOT NULL DEFAULT 'free',
            created    TEXT NOT NULL
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS payments (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id    INTEGER NOT NULL,
            plan       TEXT NOT NULL,
            amount     REAL NOT NULL,
            method     TEXT DEFAULT 'card',
            txn_id     TEXT,
            status     TEXT DEFAULT 'pending',
            note       TEXT,
            created    TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS payment_requests (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id    INTEGER NOT NULL,
            plan       TEXT NOT NULL,
            amount     REAL NOT NULL,
            method     TEXT NOT NULL,
            txn_id     TEXT,
            status     TEXT DEFAULT 'pending',
            created    TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )""")
        conn.commit()

init_db()

def hash_pw(pw):
    return hashlib.sha256(("nm2025:" + pw).encode()).hexdigest()

def get_current_user():
    uid = session.get("user_id")
    if not uid:
        return None
    try:
        with get_db() as conn:
            u = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
        return dict(u) if u else None
    except:
        return None

def check_plan(feature):
    u = get_current_user()
    if not u: return False
    return PLANS.get(u.get("plan","free"), PLANS["free"]).get(feature, False)

def groq_chat(messages, model="llama-3.3-70b-versatile", temperature=0.7, max_tokens=2048):
    r = req.post(GROQ_URL,
        headers={"Authorization":f"Bearer {GROQ_API_KEY}","Content-Type":"application/json"},
        json={"model":model,"messages":messages,"temperature":temperature,"max_tokens":max_tokens},
        timeout=30)
    r.raise_for_status()
    d = r.json()
    return d["choices"][0]["message"]["content"], d.get("usage",{}).get("total_tokens",0)

# ══════════════════════════════════════════════
#  AUTH
# ══════════════════════════════════════════════
@app.route("/login")
def login_page():
    if session.get("user_id"): return redirect("/")
    return render_template("login.html")

@app.route("/register")
def register_page():
    if session.get("user_id"): return redirect("/")
    return render_template("register.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect("/")

@app.route("/api/register", methods=["POST"])
def api_register():
    d  = request.get_json() or {}
    un = d.get("username","").strip()
    em = d.get("email","").strip().lower()
    pw = d.get("password","").strip()

    if not un or not em or not pw:
        return jsonify({"error":"Barcha maydonlarni to'ldiring"}), 400
    if len(un) < 3:
        return jsonify({"error":"Username kamida 3 ta harf"}), 400
    if len(pw) < 6:
        return jsonify({"error":"Parol kamida 6 ta belgi"}), 400
    if "@" not in em or "." not in em:
        return jsonify({"error":"Email noto'g'ri"}), 400

    try:
        with get_db() as conn:
            conn.execute(
                "INSERT INTO users (username,email,password,plan,created) VALUES (?,?,?,?,?)",
                (un, em, hash_pw(pw), "free", datetime.now().isoformat())
            )
            conn.commit()
            u = conn.execute("SELECT * FROM users WHERE email=?", (em,)).fetchone()
        session.permanent = True
        session["user_id"]  = u["id"]
        session["username"] = u["username"]
        return jsonify({"ok":True,"username":un})
    except sqlite3.IntegrityError as e:
        err = str(e).lower()
        if "username" in err:
            return jsonify({"error":"Bu username band"}), 400
        return jsonify({"error":"Bu email allaqachon ro'yxatdan o'tgan"}), 400

@app.route("/api/login", methods=["POST"])
def api_login():
    d  = request.get_json() or {}
    em = d.get("email","").strip().lower()
    pw = d.get("password","").strip()

    if not em or not pw:
        return jsonify({"error":"Email va parolni kiriting"}), 400

    with get_db() as conn:
        u = conn.execute(
            "SELECT * FROM users WHERE LOWER(email)=? AND password=?",
            (em, hash_pw(pw))
        ).fetchone()

    if not u:
        return jsonify({"error":"Email yoki parol noto'g'ri"}), 401

    session.permanent = True
    session["user_id"]  = u["id"]
    session["username"] = u["username"]
    return jsonify({"ok":True,"username":u["username"],"plan":u["plan"]})

@app.route("/api/me")
def api_me():
    u = get_current_user()
    if not u: return jsonify({"logged_in":False})
    return jsonify({
        "logged_in":True,
        "username":u["username"],
        "email":u["email"],
        "plan":u["plan"],
        "plan_info":PLANS.get(u["plan"],PLANS["free"])
    })

# ══════════════════════════════════════════════
#  PAYMENT — To'lov so'rovi (PayMe / Click / Bank)
# ══════════════════════════════════════════════

@app.route("/api/submit-payment", methods=["POST"])
def submit_payment():
    """Foydalanuvchi to'lov ma'lumotlarini yuboradi — admin tasdiqlaydi"""
    if not session.get("user_id"):
        return jsonify({"error":"Avval tizimga kiring"}), 401

    d      = request.get_json() or {}
    plan   = d.get("plan","")
    method = d.get("method","")   # payme | click | bank
    txn_id = d.get("txn_id","").strip()

    if plan not in ["pro","premium","ultra"]:
        return jsonify({"error":"Noto'g'ri plan"}), 400
    if method not in ["payme","click","bank"]:
        return jsonify({"error":"To'lov usulini tanlang"}), 400
    if not txn_id:
        return jsonify({"error":"Tranzaksiya ID (chek raqami) kiriting"}), 400

    amount = PLANS[plan]["price"]
    uid    = session["user_id"]

    with get_db() as conn:
        # Duplikat tekshirish
        existing = conn.execute(
            "SELECT id FROM payment_requests WHERE txn_id=? AND status='pending'",
            (txn_id,)
        ).fetchone()
        if existing:
            return jsonify({"error":"Bu chek raqami allaqachon yuborilgan"}), 400

        conn.execute(
            "INSERT INTO payment_requests (user_id,plan,amount,method,txn_id,status,created) VALUES (?,?,?,?,?,?,?)",
            (uid, plan, amount, method, txn_id, "pending", datetime.now().isoformat())
        )
        conn.commit()

    return jsonify({"ok":True,"message":"To'lov so'rovi yuborildi! Admin 1-24 soat ichida tasdiqlaydi."})


@app.route("/api/admin/confirm-payment", methods=["POST"])
def admin_confirm():
    """Admin to'lovni tasdiqlaydi"""
    d = request.get_json() or {}
    if d.get("secret") != ADMIN_SECRET:
        return jsonify({"error":"Ruxsat yo'q"}), 403

    req_id = d.get("request_id")
    with get_db() as conn:
        pr = conn.execute("SELECT * FROM payment_requests WHERE id=?", (req_id,)).fetchone()
        if not pr:
            return jsonify({"error":"So'rov topilmadi"}), 404
        conn.execute("UPDATE payment_requests SET status='paid' WHERE id=?", (req_id,))
        conn.execute("UPDATE users SET plan=? WHERE id=?", (pr["plan"], pr["user_id"]))
        conn.execute(
            "INSERT INTO payments (user_id,plan,amount,method,txn_id,status,created) VALUES (?,?,?,?,?,?,?)",
            (pr["user_id"],pr["plan"],pr["amount"],pr["method"],pr["txn_id"],"paid",datetime.now().isoformat())
        )
        conn.commit()
    return jsonify({"ok":True})


@app.route("/api/admin/pending-payments")
def admin_pending():
    d = request.args
    if d.get("secret") != ADMIN_SECRET:
        return jsonify({"error":"Ruxsat yo'q"}), 403
    with get_db() as conn:
        rows = conn.execute("""
            SELECT pr.id, pr.plan, pr.amount, pr.method, pr.txn_id, pr.created,
                   u.username, u.email
            FROM payment_requests pr
            JOIN users u ON u.id = pr.user_id
            WHERE pr.status='pending'
            ORDER BY pr.created DESC
        """).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/admin/payments")
def admin_payments_page():
    secret = request.args.get("secret","")
    if secret != ADMIN_SECRET:
        return "Ruxsat yo'q", 403
    with get_db() as conn:
        rows = conn.execute("""
            SELECT pr.id, pr.plan, pr.amount, pr.method, pr.txn_id, pr.created, pr.status,
                   u.username, u.email
            FROM payment_requests pr
            JOIN users u ON u.id = pr.user_id
            ORDER BY pr.created DESC LIMIT 100
        """).fetchall()
    return render_template("admin_payments.html", payments=[dict(r) for r in rows], secret=secret)


# ══════════════════════════════════════════════
#  PAGES
# ══════════════════════════════════════════════
@app.route("/")
def home():
    return render_template("index.html", user=get_current_user())

@app.route("/plans")
def plans_page():
    u = get_current_user()
    return render_template("plans.html", user=u, plans=PLANS,
                           payme_id=PAYME_ID, click_id=CLICK_ID)

@app.route("/about")
def about():
    return render_template("about.html", user=get_current_user())

@app.route("/contact")
def contact():
    return render_template("contact.html", user=get_current_user())

@app.route("/usage")
def usage():
    return render_template("usage.html", user=get_current_user())

@app.route("/creator")
def creator():
    return render_template("creator.html", user=get_current_user())

@app.route("/ai")
def ai():
    u = get_current_user()
    accessible = MODELS if (u and check_plan("models")) else {FREE_MODEL: MODELS[FREE_MODEL]}
    return render_template("ai.html", models=accessible, all_models=MODELS, user=u, plans=PLANS)

@app.route("/imagine")
def imagine_page():
    u = get_current_user()
    return render_template("imagine.html", user=u, has_access=u and check_plan("imagine"), plans=PLANS)

@app.route("/slides")
def slides_page():
    u = get_current_user()
    return render_template("slides.html", user=u, has_access=u and check_plan("slides"), plans=PLANS)

# ══════════════════════════════════════════════
#  CHAT API
# ══════════════════════════════════════════════
@app.route("/api/chat", methods=["POST"])
def chat():
    d = request.get_json() or {}
    msg   = d.get("message","").strip()
    hist  = d.get("history",[])
    mid   = d.get("model", FREE_MODEL)
    temp  = float(d.get("temperature",0.7))
    tone  = d.get("tone","default")
    if not msg: return jsonify({"error":"Bo'sh xabar"}), 400
    u = get_current_user()
    if mid not in MODELS: mid = FREE_MODEL
    if not (u and check_plan("models")): mid = FREE_MODEL
    tones = {
        "default":"Har doim foydali, do'stona va aniq javob bering.",
        "formal":"Rasmiy va professional uslubda javob bering.",
        "creative":"Ijodiy, qiziqarli va original uslubda javob bering.",
        "simple":"Juda oddiy, qisqa va tushunarli tilda javob bering.",
        "humorous":"Hazilkash va quvnoq uslubda javob bering.",
    }
    sys_p = ("Siz NovaMind AI — Tolibov Muhammad tomonidan yaratilgan aqlli yordamchi.\n"
             "Siz O'zbek, Rus va Ingliz tillarida muloqot qila olasiz.\n"
             + tones.get(tone,tones["default"])
             + "\nMarkdown formatidan foydalaning: **qalin**, *kursiv*, `kod`, ```blok```, # sarlavha.")
    msgs = [{"role":"system","content":sys_p}]
    for m in hist[-20:]: msgs.append({"role":m["role"],"content":m["content"]})
    msgs.append({"role":"user","content":msg})
    try:
        reply, tokens = groq_chat(msgs, model=mid, temperature=temp)
        return jsonify({"reply":reply,"tokens":tokens,"model":mid})
    except Exception as e:
        return jsonify({"reply":f"⚠️ Xatolik: {str(e)}","error":True}), 500

@app.route("/api/slides", methods=["POST"])
def create_slides():
    u = get_current_user()
    if not (u and check_plan("slides")):
        return jsonify({"error":"Bu funksiya faqat Ultra plan uchun!","upgrade":True}), 403
    d = request.get_json() or {}
    topic = d.get("topic","").strip()
    lang  = d.get("lang","uz")
    count = min(int(d.get("count",8)),15)
    style = d.get("style","professional")
    if not topic: return jsonify({"error":"Mavzu kiritilmadi"}), 400
    lang_map = {"uz":"O'zbek tilida","ru":"Rus tilida","en":"English"}
    sys_p = 'Siz professional prezentatsiya yaratuvchi AI. FAQAT JSON: {"title":"...","subtitle":"...","author":"NovaMind AI","slides":[{"type":"title","title":"...","subtitle":"...","notes":"..."},{"type":"content","title":"...","points":["..."],"notes":"..."},{"type":"end","title":"...","message":"...","notes":"..."}]}'
    user_p = f"Mavzu: {topic}Soni: {count}Til: {lang_map.get(lang,'Ozbek')}Uslub: {style}\nFAQAT JSON."
    try:
        raw,tokens = groq_chat(
            [{"role":"system","content":sys_p},{"role":"user","content":user_p}],
            model="llama-3.3-70b-versatile", temperature=0.6, max_tokens=4000
        )
        raw = raw.strip()
        if raw.startswith("```"):
            raw = raw.split("```")[1]
            if raw.startswith("json"): raw = raw[4:]
        data = json.loads(raw.strip().rstrip("```").strip())
        return jsonify({"success":True,"data":data,"tokens":tokens})
    except json.JSONDecodeError as e:
        return jsonify({"error":f"JSON xatolik: {str(e)}"}), 500
    except Exception as e:
        return jsonify({"error":str(e)}), 500

@app.route("/api/vision", methods=["POST"])
def vision():
    msg  = request.form.get("message","Bu rasmda nima bor?").strip()
    file = request.files.get("image")
    if not file: return jsonify({"reply":"⚠️ Rasm yuborilmadi.","error":True}), 400
    b64  = base64.b64encode(file.read()).decode()
    mime = file.mimetype or "image/jpeg"
    try:
        r = req.post(GROQ_URL,
            headers={"Authorization":f"Bearer {GROQ_API_KEY}","Content-Type":"application/json"},
            json={"model":"meta-llama/llama-4-scout-17b-16e-instruct","messages":[
                {"role":"user","content":[
                    {"type":"image_url","image_url":{"url":f"data:{mime};base64,{b64}"}},
                    {"type":"text","text":"O'zbek/Rus/Ingliz tilida javob bering.\n\n"+msg}
                ]}
            ],"max_tokens":1024}, timeout=30)
        r.raise_for_status()
        d = r.json()
        return jsonify({"reply":d["choices"][0]["message"]["content"],
                        "tokens":d.get("usage",{}).get("total_tokens",0)})
    except Exception as e:
        return jsonify({"reply":f"⚠️ Vision xatolik: {str(e)}","error":True}), 500

@app.route("/api/pdf", methods=["POST"])
def pdf_analyze():
    msg  = request.form.get("message","Bu PDF ni tahlil qil.").strip()
    file = request.files.get("file")
    if not file: return jsonify({"reply":"⚠️ Fayl yuborilmadi.","error":True}), 400
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file.read()))
        text   = "\n".join(p.extract_text() or "" for p in reader.pages).strip()
        if not text: return jsonify({"reply":"⚠️ PDF dan matn topilmadi.","error":True}), 400
        if len(text)>12000: text=text[:12000]+"...[qisqartirildi]"
        reply,tokens = groq_chat([
            {"role":"system","content":"Siz hujjat tahlilchisisiz. O'zbek/Rus/Ingliz tilida javob bering."},
            {"role":"user","content":f"PDF ({len(reader.pages)} bet):\n\n{text}\n\n---\n{msg}"}
        ], temperature=0.4, max_tokens=2048)
        return jsonify({"reply":reply,"tokens":tokens,"pages":len(reader.pages),"file_type":"pdf"})
    except Exception as e:
        return jsonify({"reply":f"⚠️ PDF xatolik: {str(e)}","error":True}), 500

@app.route("/api/word", methods=["POST"])
def word_analyze():
    msg  = request.form.get("message","Bu Word ni tahlil qil.").strip()
    file = request.files.get("file")
    if not file: return jsonify({"reply":"⚠️ Fayl yuborilmadi.","error":True}), 400
    try:
        from docx import Document
        doc   = Document(io.BytesIO(file.read()))
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        text  = "\n".join(paras)
        if not text: return jsonify({"reply":"⚠️ Word da matn topilmadi.","error":True}), 400
        if len(text)>12000: text=text[:12000]+"...[qisqartirildi]"
        reply,tokens = groq_chat([
            {"role":"system","content":"Siz hujjat tahlilchisisiz."},
            {"role":"user","content":f"Word ({len(paras)} paragraf):\n\n{text}\n\n---\n{msg}"}
        ], temperature=0.4, max_tokens=2048)
        return jsonify({"reply":reply,"tokens":tokens,"pages":len(paras),"file_type":"word"})
    except Exception as e:
        return jsonify({"reply":f"⚠️ Word xatolik: {str(e)}","error":True}), 500

@app.route("/api/pptx", methods=["POST"])
def pptx_analyze():
    msg  = request.form.get("message","Bu prezentatsiyani tahlil qil.").strip()
    file = request.files.get("file")
    if not file: return jsonify({"reply":"⚠️ Fayl yuborilmadi.","error":True}), 400
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file.read()))
        slides = []
        for i,s in enumerate(prs.slides,1):
            txts = [sh.text.strip() for sh in s.shapes if hasattr(sh,"text") and sh.text.strip()]
            if txts: slides.append(f"[Slayd {i}]: "+" | ".join(txts))
        if not slides: return jsonify({"reply":"⚠️ Prezentatsiyada matn topilmadi.","error":True}), 400
        text = "\n".join(slides)
        if len(text)>12000: text=text[:12000]+"...[qisqartirildi]"
        reply,tokens = groq_chat([
            {"role":"system","content":"Siz prezentatsiya tahlilchisisiz."},
            {"role":"user","content":f"PowerPoint ({len(prs.slides)} slayd):\n\n{text}\n\n---\n{msg}"}
        ], temperature=0.4, max_tokens=2048)
        return jsonify({"reply":reply,"tokens":tokens,"pages":len(prs.slides),"file_type":"pptx"})
    except Exception as e:
        return jsonify({"reply":f"⚠️ PPTX xatolik: {str(e)}","error":True}), 500

@app.route("/api/excel", methods=["POST"])
def excel_analyze():
    msg  = request.form.get("message","Bu Excel ni tahlil qil.").strip()
    file = request.files.get("file")
    if not file: return jsonify({"reply":"⚠️ Fayl yuborilmadi.","error":True}), 400
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
        sheets=[]; total=0
        for sn in wb.sheetnames:
            ws=wb[sn]; rows=[]
            for row in ws.iter_rows(values_only=True):
                rv=[str(c) if c is not None else "" for c in row]
                if any(v.strip() for v in rv): rows.append(" | ".join(rv)); total+=1
                if total>200: break
            if rows: sheets.append(f"[{sn}]\n"+"\n".join(rows))
            if total>200: break
        if not sheets: return jsonify({"reply":"⚠️ Excel da ma'lumot topilmadi.","error":True}), 400
        text="\n\n".join(sheets)
        if len(text)>12000: text=text[:12000]+"...[qisqartirildi]"
        reply,tokens = groq_chat([
            {"role":"system","content":"Siz Excel tahlilchisisiz."},
            {"role":"user","content":f"Excel ({len(wb.sheetnames)} varaq):\n\n{text}\n\n---\n{msg}"}
        ], temperature=0.3, max_tokens=2048)
        return jsonify({"reply":reply,"tokens":tokens,"pages":total,"file_type":"excel"})
    except Exception as e:
        return jsonify({"reply":f"⚠️ Excel xatolik: {str(e)}","error":True}), 500

@app.route("/api/imagine", methods=["POST"])
def imagine_api():
    u = get_current_user()
    if not (u and check_plan("imagine")):
        return jsonify({"error":"Bu funksiya faqat Premium va Ultra plan uchun!","upgrade":True}), 403
    d = request.get_json() or {}
    prompt = d.get("prompt","").strip()
    style  = d.get("style","")
    width  = d.get("width",800)
    height = d.get("height",600)
    if not prompt: return jsonify({"error":"Prompt bo'sh"}), 400
    try:
        enhanced,_ = groq_chat([
            {"role":"system","content":f"Translate to English and rewrite as detailed image prompt.{' Style:'+style+'.' if style else ''} Return ONLY the prompt, max 80 words."},
            {"role":"user","content":prompt}
        ], model="llama-3.1-8b-instant", temperature=0.7, max_tokens=150)
        enhanced = enhanced.strip()
    except: enhanced = prompt
    encoded = urllib.parse.quote(enhanced)
    seed    = abs(hash(prompt+style)) % 999999
    return jsonify({"image_url":f"https://image.pollinations.ai/prompt/{encoded}?width={width}&height={height}&seed={seed}&nologo=true&enhance=true",
                    "enhanced_prompt":enhanced,"original_prompt":prompt})

@app.route("/api/models")
def get_models(): return jsonify(MODELS)

@app.route("/api/contact", methods=["POST"])
def contact_send():
    name = (request.get_json() or {}).get("name","Noma'lum")
    return jsonify({"success":True,"message":f"Rahmat, {name}! Xabaringiz qabul qilindi. 🎉"})

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
